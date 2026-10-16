"""Export filled Quote Exports columns as woven-quote drafts."""
import json
from datetime import datetime, date
from openpyxl import load_workbook

SRC = r"C:\Users\Admin\Downloads\Quote Woven Pricing R1 - Copy (8).xlsx"
OUT = r"c:\Users\Admin\Desktop\approval\PO_approvalapp\POApprovalAPI\Data\WovenQuote\quotes.json"

FABRIC = [
    (70, "PP"), (71, "HD"), (72, "LLDPE"), (73, "LDPE_FY"), (74, "LDPE7"), (75, "PP_H350"),
    (76, "LDPE4"), (77, "RP"), (78, "FILLER"), (79, "TPT"), (80, "UV"), (81, "COL"),
    (82, "BLACK"), (83, "CONDUCTIVE"),
]
COATING = [
    (86, "PP"), (87, "HD"), (88, "LLDPE"), (89, "LDPE_FY"), (90, "LDPE7"), (91, "PP_H350"),
    (92, "LDPE4"), (93, "RP"), (94, "FILLER"), (95, "TPT"), (96, "UV"), (97, "COL"), (98, "BLACK"),
]
LINER = [
    (101, "LLDPE"), (102, "LDPE_FY"), (103, "LDPE7"), (104, "PP_H350"), (105, "LDPE4"),
    (106, "RP"), (107, "FILLER"), (108, "TPT"), (109, "UV"), (110, "COL"), (111, "BLACK"),
]

wb = load_workbook(SRC, data_only=True)
ws = wb["Quote Exports"]


def text(value):
    if value is None:
        return ""
    return str(value).strip()


def num(value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        value = value.strip()
        if value == "":
            return None
        try:
            return float(value)
        except ValueError:
            return None
    if isinstance(value, bool):
        return None
    return float(value)


def day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def yes(value):
    return text(value).lower() == "yes"


def recipe(col, pairs):
    out = {}
    for row, code in pairs:
        value = num(ws.cell(row, col).value)
        if value:
            out[code] = value
    return out


quotes = []
for col in range(5, 50):
    customer = text(ws.cell(7, col).value)
    if not customer:
        continue
    quote_date = day(ws.cell(5, col).value)
    quotes.append({
        "productType": "WovenBag",
        "quoteDate": quote_date,
        "sector": text(ws.cell(6, col).value) or "Export",
        "customerName": customer,
        "ourRef": text(ws.cell(8, col).value),
        "customerRef": text(ws.cell(9, col).value),
        "bagType": text(ws.cell(10, col).value) or "Tubular Bag",
        "fabricColour": text(ws.cell(11, col).value),
        "faceWidthCm": num(ws.cell(12, col).value),
        "gussetCm": num(ws.cell(13, col).value),
        "lengthCm": num(ws.cell(14, col).value),
        "top": text(ws.cell(15, col).value),
        "totalWeightGram": num(ws.cell(17, col).value),
        "fabricGsm": num(ws.cell(18, col).value),
        "linerMicron": num(ws.cell(21, col).value),
        "glossyBoppMicron": num(ws.cell(22, col).value),
        "mattBoppMicron": num(ws.cell(23, col).value),
        "metalisedBoppMicron": num(ws.cell(24, col).value),
        "whitePaperGsm": num(ws.cell(25, col).value),
        "brownPaperGsm": num(ws.cell(26, col).value),
        "adhesiveGsm": num(ws.cell(27, col).value),
        "inkGsm": num(ws.cell(28, col).value),
        "coatingGsm": num(ws.cell(30, col).value),
        "handleGram": num(ws.cell(31, col).value),
        "fabricRecipe": recipe(col, FABRIC),
        "coatingRecipe": recipe(col, COATING),
        "linerRecipe": recipe(col, LINER),
        "fabricConversion": num(ws.cell(132, col).value),
        "coatingConversion": num(ws.cell(133, col).value),
        "linerConversion": num(ws.cell(134, col).value),
        "wastageFabric": num(ws.cell(136, col).value),
        "wastageLiner": num(ws.cell(138, col).value),
        "wastageFinishing": num(ws.cell(139, col).value),
        "realisationFabric": num(ws.cell(140, col).value),
        "realisationLiner": num(ws.cell(142, col).value),
        "bagFinishing": num(ws.cell(146, col).value),
        "thumbNotch": yes(ws.cell(187, col).value),
        "metalDetection": yes(ws.cell(188, col).value),
        "perforation": yes(ws.cell(189, col).value),
        "flexoPrinting": num(ws.cell(150, col).value),
        "rewindingPerKg": num(ws.cell(151, col).value),
        "packingPerKg": num(ws.cell(152, col).value),
        "palletPerKg": num(ws.cell(153, col).value),
        "fobCostPerKg": num(ws.cell(154, col).value),
        "paymentDays": num(ws.cell(157, col).value),
        "profitPercent": num(ws.cell(160, col).value),
        "usdInr": num(ws.cell(163, col).value),
        "euroPerUsd": num(ws.cell(164, col).value),
        "gbpPerUsd": num(ws.cell(165, col).value),
        "freightUsdPerKg": num(ws.cell(169, col).value),
        "commissionUsdPerKg": num(ws.cell(170, col).value),
        "dimensionUnit": text(ws.cell(176, col).value) or "CM",
        "terms": text(ws.cell(190, col).value) or "FOB",
        "seaPort": text(ws.cell(191, col).value),
        "currency": text(ws.cell(192, col).value) or "USD",
        "packing": text(ws.cell(194, col).value),
        "pcsPerBale": num(ws.cell(195, col).value),
        "palletSize": text(ws.cell(196, col).value),
        "bagsPerPallet": num(ws.cell(197, col).value),
        "printPlateCost": num(ws.cell(198, col).value),
        "remarks": text(ws.cell(199, col).value),
        "sheetPrice": num(ws.cell(193, col).value),
    })

print("quotes", len(quotes))
for q in quotes:
    print(q["quoteDate"], q["customerName"], q["ourRef"], q["customerRef"], "sheet", q["sheetPrice"])
with open(OUT, "w", encoding="utf-8") as handle:
    json.dump(quotes, handle, ensure_ascii=False, indent=2)
print("wrote", OUT)
