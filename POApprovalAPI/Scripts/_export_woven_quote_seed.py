"""One-off export of Quote Woven Pricing masters into seed.json."""
import json
from datetime import datetime, date
from openpyxl import load_workbook

SRC = r"C:\Users\Admin\Downloads\Quote Woven Pricing R1 - Copy (8).xlsx"
OUT = r"c:\Users\Admin\Desktop\approval\PO_approvalapp\POApprovalAPI\Data\WovenQuote\seed.json"

MATERIALS = [
    ("PP", "PP Rafia H030SG", 1, True, True, False),
    ("HD", "HD Rafia E52009", 2, True, True, False),
    ("LLDPE", "LLDPE F19010 Film", 3, True, True, True),
    ("LDPE_FY", "LDPE FY1005 Film", 4, True, True, True),
    ("LDPE7", "LDPE 7 MFI 1070LA17", 5, True, True, True),
    ("PP_H350", "PP H350FG", 6, True, True, True),
    ("LDPE4", "LDPE 4 MFI J24FS040", 7, True, True, True),
    ("RP", "RP", 8, True, True, True),
    ("FILLER", "Filler", 9, True, True, True),
    ("TPT", "TPT", 10, True, True, True),
    ("UV", "UV", 11, True, True, True),
    ("COL", "Col MB", 12, True, True, True),
    ("BLACK", "Black MB", 13, True, True, True),
    ("CONDUCTIVE", "Conductive Yarn", 14, True, False, False),
    ("BACKSEAM", "Back Seam Coating RM", 15, False, False, False),
    ("WHITE_PAPER", "White Paper", 16, False, False, False),
    ("BROWN_PAPER", "Brown Paper", 17, False, False, False),
    ("GLOSSY", "Glossy BOPP", 18, False, False, False),
    ("MATT", "Matt BOPP", 19, False, False, False),
    ("METAL", "Metalised BOPP", 20, False, False, False),
    ("ADHESIVE", "Adhesive", 21, False, False, False),
    ("INK", "Ink BOPP", 22, False, False, False),
    ("SOLVENT", "Solvant BOPP", 23, False, False, False),
]

# Row on both RM blocks, in material order (rows 5-27 and 32-54).
RATE_ROWS = list(range(5, 28))


def iso(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def num(value):
    if value is None or isinstance(value, str):
        return None
    return float(value)


wb = load_workbook(SRC, data_only=True)
rm = wb["RM Prices"]
codes = [m[0] for m in MATERIALS]

rates = []
for sector, header_row, first_data_row in (("Export", 4, 5), ("Domestic", 31, 32)):
    for col in range(3, 80):
        when = iso(rm.cell(header_row, col).value)
        if not when:
            continue
        for offset, code in enumerate(codes):
            rate = num(rm.cell(first_data_row + offset, col).value)
            if rate is None:
                continue
            rates.append({"code": code, "sector": sector, "effectiveDate": when, "rate": round(rate, 6)})

cm = wb["Cust Master"]
customers = []
seen = set()
for sector, start, end in (("Export", 3, 345), ("Domestic", 347, 530)):
    for row in range(start, end):
        full = cm.cell(row, 3).value
        short = cm.cell(row, 4).value
        if not full:
            continue
        full = str(full).strip()
        short = str(short).strip() if short else full
        key = (sector, full.lower())
        if key in seen:
            continue
        seen.add(key)
        customers.append({"sector": sector, "fullName": full, "shortName": short})

lookups = []
master_groups = [
    ("Top", ["Heat Cut", "Hemmed", "Ez Open", "No Fold"]),
    ("BagType", ["Back Seam Bag", "Tubular Bag", "Tubular Fabric", "Flat Fabric"]),
    ("Colour", ["MW", "Black/White", "Transparent", "Coloured", "MW With Coloured Stripe", "MW With Window"]),
    ("Terms", ["Ex Works", "FOB", "CIF", "DDP", "DDU"]),
    ("Currency", ["INR", "USD", "Euro", "GBP"]),
    ("DimensionUnit", ["Inch", "CM"]),
    ("Packing", ["Floor Loaded", "Box/Cartoon", "Palletised"]),
    ("Sector", ["Export", "Domestic"]),
]
for category, values in master_groups:
    for index, value in enumerate(values, start=1):
        lookups.append({"category": category, "value": value, "sortOrder": index})

constants = [
    ("density", 0.92, "Film density", "Weight"),
    ("cutHeat", 5, "Heat cut extra cm", "Weight"),
    ("cutHem", 8, "Hemmed extra cm", "Weight"),
    ("cutEz", 0, "Ez Open extra cm", "Weight"),
    ("cutNoFold", 0, "No Fold extra cm", "Weight"),
    ("backSeamExtraCm", 2.5, "Back seam width extra cm", "Weight"),
    ("linerWidthExtra", 2, "Liner width extra cm", "Weight"),
    ("linerLengthExtra", 5, "Liner length extra cm", "Weight"),
    ("boppWidthExtra", 1.5, "BOPP / paper width extra cm", "Weight"),
    ("coatingWidthExtra", 2, "Coating width extra cm", "Weight"),
    ("easyOpenGramPerMeter", 7, "Easy open grams per metre", "Weight"),
    ("easyOpenWidthExtra", 10, "Easy open width extra cm", "Weight"),
    ("threadHem", 3, "Hemmed thread factor", "Weight"),
    ("threadHeat", 2, "Heat cut thread factor", "Weight"),
    ("threadEz", 1.2, "Ez Open thread factor", "Weight"),
    ("threadDivisor", 61, "Thread divisor", "Weight"),
    ("threadEzWidthExtra", 10, "Ez Open thread width extra cm", "Weight"),
    ("backSeamCoatingGrPerM", 2, "Back seam coating g/m", "Weight"),
    ("solventInkFactor", 2, "Solvent weight multiple of ink", "Print"),
    ("printWasteFraction", 0.05, "Printed film waste fraction", "Print"),
    ("printConvPaper", 10, "Paper printing conversion Rs/kg", "Print"),
    ("printConvBopp", 40, "BOPP printing conversion Rs/kg", "Print"),
    ("coatWasteTubularNum", 1.75, "Tubular coating waste numerator", "Wastage"),
    ("coatWasteOtherNum", 3.75, "Other coating waste numerator", "Wastage"),
    ("coatWasteBase", 0.0075, "Coating waste base fraction", "Wastage"),
    ("coatingWastageUsesTotalWeight", 1, "1 = coating wastage on total weight, as in the workbook", "Wastage"),
    ("wcRate", 0.125, "Working capital rate", "Price"),
    ("wcBaseDays", 30, "Working capital base days", "Price"),
    ("yearDays", 365, "Days in the working-capital year", "Price"),
    ("fixedAssetPerKg", 13.14, "Interest and depreciation Rs/kg", "Price"),
    ("thumbNotch", 0.4, "Thumb notch Rs/pc", "Finishing"),
    ("metalDetection", 0.3, "Metal detection Rs/pc", "Finishing"),
    ("perforation", 0.2, "Perforation Rs/pc", "Finishing"),
    ("realisationCoatingNoPaper", 35, "Coating realisation Rs/kg without paper", "Wastage"),
    ("realisationCoatingPaper", 0, "Coating realisation Rs/kg with paper", "Wastage"),
    ("realisationFinishingNoPaper", 35, "Finishing realisation Rs/kg without paper", "Wastage"),
    ("realisationFinishingPaper", 8, "Finishing realisation Rs/kg with paper", "Wastage"),
    ("cifAddsCommissionAgain", 1, "1 = CIF adds commission a second time, as in the workbook", "Price"),
]

payload = {
    "materials": [
        {
            "code": code,
            "name": name,
            "sortOrder": sort,
            "inFabric": fabric,
            "inCoating": coating,
            "inLiner": liner,
        }
        for code, name, sort, fabric, coating, liner in MATERIALS
    ],
    "rates": rates,
    "customers": customers,
    "lookups": lookups,
    "constants": [
        {"key": key, "value": value, "label": label, "groupName": group}
        for key, value, label, group in constants
    ],
}
print("rates", len(rates), "customers", len(customers))
with open(OUT, "w", encoding="utf-8") as handle:
    json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
print("wrote", OUT)
